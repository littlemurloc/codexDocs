from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "deliverables" / "三国斗阵_第一批新增预选英雄技能设计归档_v1.0.docx"

FONT = "Microsoft YaHei"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2937"
MUTED = "5B6573"
HEADER_FILL = "E8EEF5"
CALLOUT_FILL = "F4F6F9"
GRID = "C9D2DE"


HEROES = [
    {
        "name": "曹仁",
        "epithet": "据守如山",
        "battle": "近战前排，攻击距离 300，带领近战兵",
        "position": "前排",
        "role": "以全军兵损转化为防守能力，建立阵地防线。",
        "core_name": "守备",
        "core": "每有 1 名我方正式部署的小兵死亡，曹仁获得 1 层守备；每 6 秒额外获得 1 层，最多 5 层。本波次内保留，波次结束清空；临时召唤物死亡不提供守备。",
        "skills": [
            ("1 星", "固阵拒敌", "主动", "自动对当前目标造成 180% 攻击的物理伤害；当前目标为敌方队长或精英时，嘲讽其 3 秒。CD = 8 秒。"),
            ("2 星", "八门守备", "被动", "启用守备：每有 1 名我方正式部署小兵死亡，曹仁 +1 层；每 6 秒额外 +1 层；最多 5 层；本波次保留、波次结束清空；临时召唤物不计入。战前前排时，受到近战伤害 -10%。"),
            ("3 星", "横戈固垒", "主动", "仅当生命低于 60% 的我方队长不超过 1 名，且当前目标区域存在敌方队长或精英时自动释放。对区域最多 3 名敌人各造成 120% 攻击物理伤害；守备至少 3 层时，命中的队长/精英攻击 -20%，持续 5 秒。不消耗守备；CD = 11 秒。"),
            ("4 星", "合围守势", "被动", "每波首次守备达到 3 层时，所有我方小兵获得“固守”：受到伤害 -10%，持续至本波结束。同一波内不会因守备变化而消失、刷新或重复触发。"),
            ("5 星", "铁壁沉舟", "主动", "当至少 2 名我方队长生命低于 60%，且守备至少 1 层时自动释放。消耗全部守备；所有我方队长和小兵获得自身最大生命 × 5% × 消耗层数的护盾，持续 7 秒；CD = 15 秒。"),
            ("6 星", "死守不退", "被动", "每波首次守备达到 5 层时自动触发，不消耗守备。所有我方队长受到伤害 -10%，所有我方小兵受到伤害 -15%，持续 6 秒；每波 1 次。"),
        ],
        "readability": "守备显示为 0-5 层城垛；小兵阵亡与每次 6 秒自然获取均显示“守备 +1”。第 4 星固守、第 5 星消耗层数及护盾值、第 6 星全军减伤均须独立提示。",
    },
    {
        "name": "典韦",
        "epithet": "恶来死战",
        "battle": "近战前排，攻击距离 300，带领近战兵",
        "position": "前排",
        "role": "残血越战越强的单点攻坚战士。",
        "core_name": "死战",
        "core": "典韦每波首次降至生命 75%、50%、25% 或以下时，各获得 1 层死战，最多 3 层；一次伤害跨越多个尚未触发的阈值时，获得对应层数。不会因治疗失去，波次结束清空。每层死战使典韦攻击提高 5%。",
        "skills": [
            ("1 星", "双戟断首", "主动", "自动对当前目标造成 190% 攻击的物理伤害；当前目标为敌方队长或精英时，攻击 -15%，持续 4 秒。CD = 8 秒。"),
            ("2 星", "古之恶来", "被动", "启用死战：生命首次降至 75%、50%、25% 或以下时各 +1 层，最多 3 层；单次跨越多个未触发阈值时获得对应层数；每层攻击 +5%；不因治疗失去，波次结束清空。战前前排时，受到近战伤害 -8%。"),
            ("3 星", "旋戟披靡", "主动", "当前目标区域至少有 2 名敌人时自动释放。对区域最多 3 名敌人各造成 120% 攻击物理伤害；死战至少 2 层时，每名命中敌人额外受 60% 攻击物理伤害。不消耗死战；CD = 11 秒。"),
            ("4 星", "浴血难退", "被动", "每层死战使典韦受到控制效果的持续时间 -10%，最多 -30%。"),
            ("5 星", "恶来绝命", "主动", "当典韦生命低于 60%、死战至少 1 层且存在敌方队长或精英时自动释放。选择生命比例最低的敌方队长或精英，造成 240% 攻击 + 每层死战 45% 攻击的物理伤害；消耗全部死战；CD = 15 秒。"),
            ("6 星", "死战不屈", "被动", "每波首次获得第 3 层死战时自动触发：清除自身 1 个普通减益，受到伤害 -15%，持续 6 秒；每波 1 次。"),
        ],
        "readability": "死战为 0-3 层染血戟图标；75% / 50% / 25% 阈值均须显示获得提示，明确它不是受击次数资源。",
    },
    {
        "name": "张辽",
        "epithet": "威震逍遥",
        "battle": "近战中前排/侧翼，攻击距离 300，带领近战兵",
        "position": "中前排或侧翼",
        "role": "标记并压制敌方关键队长或精英。",
        "core_name": "震慑",
        "core": "主动技能以敌方队长或精英为主目标时，对该目标施加 1 层震慑，最多 3 层，持续 8 秒并刷新；范围附带命中不施加、刷新或转移震慑。张辽同一时间只会维持 1 个震慑目标；对新目标施加震慑会清除旧目标的震慑，震慑目标死亡时立即消失。每层震慑使目标对张辽以外我方单位的伤害 -4%，最多 -12%；波次结束清空。",
        "skills": [
            ("1 星", "逍遥突斩", "主动", "自动对当前目标造成 190% 攻击的物理伤害。CD = 8 秒。"),
            ("2 星", "威震合肥", "被动", "启用震慑：主动以敌方队长或精英为主目标时施加 1 层，最多 3 层，持续 8 秒、重复刷新；仅维持 1 个震慑目标；范围附带命中不改变震慑；每层使该目标对张辽以外我方单位伤害 -4%，最多 -12%；波次结束清空。战前侧翼时，攻击 +10%、受到远程伤害 -10%。"),
            ("3 星", "断后破军", "主动", "震慑目标所在区域至少有 2 名敌人时自动释放。对震慑目标及区域内最多另外 2 名敌人各造成 120% 攻击物理伤害；震慑目标额外受到 80% 攻击物理伤害，并获得 1 层震慑。CD = 11 秒。"),
            ("4 星", "声威夺气", "被动", "每波中，同一敌方队长或精英首次达到 3 层震慑时，物理防御 -12，持续 6 秒；该“胆寒”状态不因后续消耗震慑而提前消失，每个目标每波 1 次。"),
            ("5 星", "八百破阵", "主动", "震慑目标拥有 3 层震慑时自动释放。对该目标造成 260% 攻击 + 每层震慑 45% 攻击的物理伤害；消耗其全部震慑；满层时为 395% 攻击；CD = 15 秒。"),
            ("6 星", "止啼之威", "被动", "每次张辽使用“八百破阵”消耗 3 层震慑后触发：张辽攻击 +15%，本队近战小兵对敌方队长或精英造成伤害 +15%，持续 6 秒。"),
        ],
        "readability": "震慑目标必须显示专属旗标与 1-3 层计数；第 4 星胆寒、第 5 星消耗层数和第 6 星小队强化均须独立提示。",
    },
    {
        "name": "马超",
        "epithet": "西凉突骑",
        "battle": "近战侧翼，攻击距离 300，带领近战兵",
        "position": "侧翼",
        "role": "通过击破普通敌人积累冲势的连续突破战士。",
        "core_name": "冲势",
        "core": "每波开始时，马超获得 1 层冲势；每有 1 名敌方正式部署的小兵死亡，马超获得 1 层冲势，每 2 秒最多因此获得 1 层；最多 3 层。波次结束清空；临时召唤物死亡不提供冲势。每层冲势使马超攻击 +4%，持有 3 层时额外攻击速度 +10%。",
        "skills": [
            ("1 星", "铁骑突袭", "主动", "自动对当前目标造成 195% 攻击物理伤害；目标为敌方队长或精英时，额外造成 45% 攻击物理伤害；CD = 8 秒。"),
            ("2 星", "西凉战马", "被动", "启用冲势：每波开始 +1 层；每有 1 名敌方正式部署小兵死亡 +1 层，每 2 秒最多 1 层；最多 3 层；波次结束清空，临时召唤物不计入。每层攻击 +4%，3 层时额外攻击速度 +10%。战前侧翼时，攻击速度 +15%。"),
            ("3 星", "长枪贯阵", "主动", "当前目标所在一路至少有 2 名敌人时自动释放。从当前目标向该路后方最多命中 4 名敌人，各造成 115% 攻击物理伤害；冲势至少 2 层时，每名命中敌人额外受 45% 攻击物理伤害。不消耗冲势；CD = 11 秒。"),
            ("4 星", "马踏连营", "被动", "每波首次冲势达到 3 层时自动触发：马超攻击速度 +15%，本队近战小兵攻击速度 +10%，持续 6 秒；每波 1 次。"),
            ("5 星", "狮盔绝尘", "主动", "冲势至少 1 层且存在敌方队长或精英时自动释放。选择生命比例最低的敌方队长或精英，造成 250% 攻击 + 每层冲势 45% 攻击的物理伤害；消耗全部冲势；CD = 15 秒。"),
            ("6 星", "破军余威", "被动", "每次马超使用“狮盔绝尘”消耗 3 层冲势后触发：目标物理防御 -12，本队近战小兵对该目标造成伤害 +15%，均持续 6 秒；不改变小兵目标。"),
        ],
        "readability": "冲势为 0-3 层马蹄图标；敌方正式小兵死亡时显示“冲势 +1”。第 4 星加速、第 5 星消耗层数、第 6 星破甲和小兵强化均须独立提示。",
    },
    {
        "name": "荀彧",
        "epithet": "王佐定策",
        "battle": "远程后排，攻击距离 600，带领远程兵",
        "position": "后排",
        "role": "用全队存活状态维持战线的策略辅助。",
        "core_name": "政令",
        "core": "每波开始时荀彧获得 2 层政令，此后每 8 秒额外获得 1 层，最多 3 层；本波次内保留，波次结束清空。政令本身不直接提供属性，只被技能用于强化或消耗。",
        "skills": [
            ("1 星", "驱虎制衡", "主动", "自动选择攻击最高的敌方队长或精英；不存在时攻击当前目标。造成 175% 攻击法术伤害，并使其造成伤害 -15%，持续 5 秒；CD = 8 秒。"),
            ("2 星", "王佐之才", "被动", "启用政令：每波开始 2 层，此后每 8 秒 +1 层，最多 3 层；波次结束清空。战前后排时，受到远程伤害 -12%。"),
            ("3 星", "奉常安阵", "主动", "至少 1 名我方队长生命低于 70% 时自动释放。为生命比例最低的至多 2 名我方队长各施加 12% 最大生命护盾，持续 6 秒；政令至少 2 层时，额外受到伤害 -10%，持续 5 秒。不消耗政令；CD = 11 秒。"),
            ("4 星", "令行禁止", "被动", "每波首次政令达到 3 层时自动触发：所有我方队长受到伤害 -8%，持续 6 秒；每波 1 次，不因后续政令变化而改变。"),
            ("5 星", "举国同心", "主动", "政令拥有 3 层且存在敌方队长或精英时自动释放。消耗全部政令；所有我方队长攻击 +15%，并清除各自 1 个普通减益，持续 7 秒；CD = 15 秒。"),
            ("6 星", "荐贤辅弼", "被动", "每次荀彧使用“举国同心”消耗 3 层政令后触发：所有我方小兵攻击速度 +15%，对敌方队长或精英造成伤害 +10%，持续 6 秒；不改变小兵目标。"),
        ],
        "readability": "政令显示为 0-3 枚令牌；第 4 星全军减伤、第 5 星全军指令与第 6 星小兵强化均须独立提示。",
    },
    {
        "name": "太史慈",
        "epithet": "独胆双戟",
        "battle": "近战前中排/侧翼，攻击距离 300，带领近战兵",
        "position": "前中排或侧翼",
        "role": "对敌方队长或精英发起单目标决斗压制。",
        "core_name": "战书",
        "core": "太史慈的主动技能以敌方队长或精英为主目标时，对其施加战书，持续 8 秒；重复以该目标为主目标时刷新。仅维持 1 个目标；施加到新目标时清除原目标，目标死亡时清除，波次结束清空；范围附带命中不改变战书。太史慈对战书目标的普攻与主动伤害 +12%。",
        "skills": [
            ("1 星", "双戟请战", "主动", "自动选择攻击范围内生命比例最低的敌方队长或精英；不存在时攻击当前目标。造成 190% 攻击物理伤害；CD = 8 秒。"),
            ("2 星", "信义赴约", "被动", "启用战书：主动以敌方队长或精英为主目标时施加，持续 8 秒、重复刷新；仅维持 1 个目标；转移、目标死亡或波次结束时清除；范围附带命中不改变战书。对战书目标普攻与主动伤害 +12%。战前侧翼时，攻击速度 +10%、受到远程伤害 -8%。"),
            ("3 星", "交锋破甲", "主动", "战书目标存在时自动释放。对其造成 170% 攻击物理伤害，物理防御 -12，持续 5 秒，并刷新战书；CD = 11 秒。"),
            ("4 星", "援军同击", "被动", "每次对新的敌方队长或精英首次施加战书时，目标获得“合击”：本队近战小兵对其自然攻击伤害 +15%，持续 6 秒。同一目标每波 1 次；“合击”不因战书消耗或转移而提前消失。"),
            ("5 星", "弑将决胜", "主动", "战书目标生命低于 50% 时自动释放。造成 300% 攻击物理伤害；目标生命低于 40% 时额外造成 110% 攻击物理伤害；随后消耗战书；CD = 15 秒。"),
            ("6 星", "信义死斗", "被动", "对同一名敌方队长或精英维持战书连续 6 秒时自动触发：清除太史慈自身 1 个普通减益，攻击 +15%、受到伤害 -12%，持续 6 秒。同一战书目标每波 1 次；战书转移、清除或目标死亡时连续时间重置。"),
        ],
        "readability": "唯一战书目标显示红色战书图标与剩余时间；合击、第 5 星消耗战书与第 6 星连续维持进度均须独立提示。",
    },
    {
        "name": "孙尚香",
        "epithet": "三才游剑",
        "battle": "远程后排，攻击距离 600，带领远程兵；武器表现为远程飞剑",
        "position": "后排",
        "role": "以三才游剑压低多目标血线；飞剑集中时单点爆发，分散时形成剑阵压制。",
        "core_name": "三才游剑",
        "core": "每波开始时，孙尚香拥有 3 柄待命游剑。飞剑普攻命中时，若有待命游剑则向目标派出 1 柄，每 1.5 秒最多 1 柄；场上最多 3 柄悬剑，同一目标可以同时悬停多柄。悬剑持续 6 秒，持续结束、目标死亡或被技能召回时回到待命状态。",
        "skills": [
            ("1 星", "流星飞剑", "主动", "自动对当前目标造成 180% 攻击物理伤害；CD = 8 秒。"),
            ("2 星", "三才御剑", "被动", "启用三才游剑：每波开始拥有 3 柄待命游剑；飞剑普攻命中时，若有待命游剑则向目标派出 1 柄，每 1.5 秒最多 1 柄；场上最多 3 柄悬剑，同一目标可同时悬停多柄；悬剑持续 6 秒，持续结束、目标死亡或被技能召回时回到待命。战前后排时，受到远程伤害 -12%。"),
            ("3 星", "剑影回环", "主动", "场上至少有 1 柄悬剑时自动释放。每柄悬剑分别对其悬停目标造成 90% 攻击物理伤害；不召回悬剑；CD = 11 秒。"),
            ("4 星", "双锋裂甲", "被动", "每波中，同一敌方单位首次同时被至少 2 柄悬剑悬停时，获得“裂甲”：物理防御 -10，持续 6 秒；不因悬剑返回、召回或数量变化而提前消失；同一目标每波 1 次。"),
            ("5 星", "三剑归一", "主动", "场上 3 柄悬剑全部悬停于同一名敌方队长或精英时自动释放。召回全部悬剑；每柄对该目标造成 125% 攻击物理伤害，合计 375%；CD = 15 秒。"),
            ("6 星", "三才乱阵", "被动", "每波首次场上同时存在 3 柄悬剑，且分布于至少 2 名不同敌人时自动触发。每柄悬剑对其悬停目标额外造成 110% 攻击物理伤害；不召回悬剑；每波 1 次。"),
        ],
        "readability": "敌人头顶显示 1-3 柄悬剑；裂甲、第 5 星集中召回与第 6 星分散剑阵均须独立提示。",
    },
    {
        "name": "夏侯渊",
        "epithet": "神速追猎",
        "battle": "远程中后排，攻击距离 600，带领远程兵",
        "position": "中排或后排",
        "role": "清理小兵后获得斩首窗口的远程追猎型输出。",
        "core_name": "追猎",
        "core": "普攻首次命中未被锁定的敌人时施加猎印，持续 4 秒；场上仅维持 1 个猎印，后续普攻不刷新。猎印自然结束或其目标死亡时，完成一次追猎：获得 1 层猎机并清除猎印，每 1.5 秒最多 1 次。猎机最多 3 层，本波次内保留、波次结束清空；每层使夏侯渊对敌方队长或精英伤害 +5%。",
        "skills": [
            ("1 星", "疾袭定标", "主动", "自动对当前目标造成 185% 攻击物理伤害；CD = 8 秒。"),
            ("2 星", "神速追猎", "被动", "启用追猎：普攻首次命中未被锁定敌人时施加猎印，持续 4 秒，仅维持 1 个且后续普攻不刷新；自然结束或目标死亡时完成追猎，获得 1 层猎机并清除猎印，每 1.5 秒最多 1 次；猎机最多 3 层，本波次保留、波次结束清空；每层对敌方队长/精英伤害 +5%。战前中排或后排时，攻击速度 +12%。"),
            ("3 星", "急袭贯阵", "主动", "存在猎印目标且该目标所在一路至少有 2 名敌人时自动释放。对猎印目标及该路最多另外 3 名敌人各造成 115% 攻击物理伤害；猎印目标额外受 80% 攻击物理伤害；不消耗猎机；CD = 11 秒。"),
            ("4 星", "乘隙疾行", "被动", "每波首次猎机达到 3 层时自动触发：夏侯渊攻击速度 +15%，本队远程小兵攻击速度 +10%，持续 6 秒；每波 1 次。"),
            ("5 星", "追魂一击", "主动", "猎机至少 1 层且存在敌方队长或精英时自动释放。选择生命比例最低的敌方队长或精英，造成 250% 攻击 + 每层猎机 45% 攻击的物理伤害；消耗全部猎机；CD = 15 秒。"),
            ("6 星", "追迹识破", "被动", "每波中，同一敌方队长或精英被猎印锁定满 4 秒且仍存活时，获得“识破”：夏侯渊对其伤害 +20%，本队远程小兵对其自然攻击伤害 +10%，持续 6 秒；同一目标每波 1 次。猎印正常结束并提供猎机；识破不因猎印结束而提前消失。"),
        ],
        "readability": "猎印显示在唯一锁定目标头顶并显示 4 秒倒计时；猎机显示为 0-3 层；识破、第 5 星消耗层数及伤害均须独立提示。",
    },
]


def set_run_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def set_table_borders(table, color=GRID):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_text(paragraph, text, *, size=10.5, color=INK, bold=False, italic=False):
    run = paragraph.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return run


def set_paragraph(paragraph, *, before=0, after=6, line=1.25, alignment=None, keep_with_next=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep_with_next
    if alignment is not None:
        paragraph.alignment = alignment


def clear_paragraph(paragraph):
    p = paragraph._element
    for child in list(p):
        p.remove(child)


def add_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    clear_paragraph(p)
    set_paragraph(p, after=0)
    add_text(p, "三国斗阵｜英雄技能设计归档", size=9, color=MUTED)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    footer = section.footer
    p = footer.paragraphs[0]
    clear_paragraph(p)
    set_paragraph(p, after=0)
    add_text(p, "内部设计归档 · 第一批新增预选英雄 · v1.0", size=8.5, color=MUTED)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    add_header_footer(section)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25
        style.paragraph_format.keep_with_next = True


def add_label_detail_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    set_table_borders(table)
    headers = ("项目", "说明")
    for index, label in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade(cell, HEADER_FILL)
        p = cell.paragraphs[0]
        set_paragraph(p, after=0, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        add_text(p, label, size=9.5, color=DARK_BLUE, bold=True)
    set_repeat_header(table.rows[0])
    for label, detail in rows:
        cells = table.add_row().cells
        shade(cells[0], HEADER_FILL)
        p = cells[0].paragraphs[0]
        set_paragraph(p, after=0)
        add_text(p, label, size=9.5, color=DARK_BLUE, bold=True)
        p = cells[1].paragraphs[0]
        set_paragraph(p, after=0)
        add_text(p, detail, size=9.5)
    set_table_geometry(table, [2700, 6660])
    return table


def add_callout(doc, label, text):
    p = doc.add_paragraph()
    set_paragraph(p, before=4, after=7)
    p.paragraph_format.left_indent = Pt(5)
    p.paragraph_format.right_indent = Pt(5)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), CALLOUT_FILL)
    p_pr.append(shd)
    add_text(p, f"{label}：", size=10, color=DARK_BLUE, bold=True)
    add_text(p, text, size=10)
    return p


def add_skill_table(doc, skills):
    table = doc.add_table(rows=1, cols=3)
    set_table_borders(table)
    headers = ("星级", "技能", "正式效果与自动规则")
    for index, label in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade(cell, HEADER_FILL)
        p = cell.paragraphs[0]
        set_paragraph(p, after=0, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        add_text(p, label, size=9.5, color=DARK_BLUE, bold=True)
    set_repeat_header(table.rows[0])
    for star, name, skill_type, effect in skills:
        cells = table.add_row().cells
        p = cells[0].paragraphs[0]
        set_paragraph(p, after=0, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        add_text(p, star, size=9.5, color=DARK_BLUE, bold=True)
        p = cells[1].paragraphs[0]
        set_paragraph(p, after=0)
        add_text(p, name, size=9.5, color=INK, bold=True)
        p.add_run("\n")
        add_text(p, skill_type, size=8.5, color=MUTED)
        p = cells[2].paragraphs[0]
        set_paragraph(p, after=0)
        add_text(p, effect, size=9.2)
    set_table_geometry(table, [900, 2000, 6460])
    return table


def add_batch_matrix(doc):
    rows = [
        ("曹仁", "全军正式小兵阵亡 + 6 秒自然获取", "防御嘲讽 / 低压区域压制 / 多队长承压时全军护盾", "与张飞、夏侯惇的承伤攒层不同：资源来自全军兵损，保留。"),
        ("典韦", "自身首次跨越 75% / 50% / 25% 生命阈值", "单点削攻 / 区域攻坚 / 低血时斩首", "以残血跨阈值推进，不与反击坦克重复，保留。"),
        ("张辽", "唯一队长/精英震慑标记", "单点 / 标记区域压制 / 满层破阵", "与太史慈同属单目标标记；分别偏敌方输出压制与决斗锁杀，列入对照实测。"),
        ("马超", "波次初始 1 层 + 敌方正式小兵死亡", "单点突破 / 一路贯穿 / 攒冲势后斩首", "与夏侯渊都存在积累后斩首；分别偏兵线突破与远程追猎，列入对照实测。"),
        ("荀彧", "波次初始 2 层 + 每 8 秒政令", "关键目标减伤 / 低血队长护盾 / 全队增攻净化", "策略辅助与刘备的低血急救区隔明确，保留。"),
        ("太史慈", "唯一队长/精英战书，连续维持", "锁定请战 / 单目标破甲 / 低血决胜", "与张辽同属标记体系；以自身及本队对决斗目标的收益区隔，列入对照实测。"),
        ("孙尚香", "3 柄游剑的待命、悬停、集中和分散", "单点飞剑 / 悬剑群体回环 / 三剑集中爆发", "场上分布状态独立，未与既有远程英雄重复，保留。"),
        ("夏侯渊", "唯一猎印自然结束或目标死亡后获得猎机", "单点定标 / 猎印一路贯阵 / 猎机斩首", "与马超同有积累后斩首；触发来自追猎完成而非兵线击杀，列入对照实测。"),
    ]
    table = doc.add_table(rows=1, cols=4)
    set_table_borders(table)
    headers = ("英雄", "核心来源", "1 / 3 / 5 星主动职责", "自检结论")
    for index, label in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade(cell, HEADER_FILL)
        p = cell.paragraphs[0]
        set_paragraph(p, after=0, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        add_text(p, label, size=8.8, color=DARK_BLUE, bold=True)
    set_repeat_header(table.rows[0])
    for row in rows:
        cells = table.add_row().cells
        for index, text in enumerate(row):
            p = cells[index].paragraphs[0]
            set_paragraph(p, after=0)
            add_text(p, text, size=8.5)
    set_table_geometry(table, [800, 2150, 3750, 2660])
    return table


def add_history_comparison_table(doc):
    rows = [
        ("已完成英雄范围", "黄忠、刘备、吕蒙、赵云、张飞、夏侯惇、吕布、诸葛亮、司马懿、甘宁。"),
        ("本批已形成的差异", "全军兵损防御、生命阈值狂战、唯一压制/决斗标记、游剑分布、猎印自然完成等机制与早期英雄的后排蓄势、隐身背刺、闪避救援、反击、施法循环不同。"),
        ("当前库的缺口", "本批仍偏重伤害、护甲、攻击和攻速窗口；位移、反击、隐身、稳定控制、战场规则等动词主要仍由早期英雄承担。下一批立项时应优先检查这些覆盖，而非继续复制资源消耗斩首。"),
        ("历史例外", "张辽、马超、荀彧的 6 星触发绑定 5 星消耗满层资源；已批准、保留，但不可复用为后续默认模板。早期英雄中含旧版参与击败或隐式优先级的表述，也不得直接复制到新设计。"),
    ]
    return add_label_detail_table(doc, rows)


def add_overview(doc):
    p = doc.add_paragraph()
    set_paragraph(p, after=3)
    add_text(p, "三国斗阵", size=12, color=DARK_BLUE, bold=True)

    p = doc.add_paragraph()
    set_paragraph(p, after=4)
    add_text(p, "第一批新增预选英雄", size=26, color="0B2545", bold=True)

    p = doc.add_paragraph()
    set_paragraph(p, after=18)
    add_text(p, "技能设计归档 v1.0", size=15, color=MUTED)

    add_label_detail_table(
        doc,
        [
            ("归档范围", "曹仁、典韦、张辽、马超、荀彧、太史慈、孙尚香、夏侯渊。"),
            ("审核状态", "8 名英雄均已通过技能审核；本文件为人类阅读版归档，后续进入统一落表。"),
            ("设计日期", "2026-08-05"),
            ("适用阶段", "P0 英雄技能设计；所有技能等级固定为 skill_rank = 1。"),
        ],
    )

    doc.add_heading("使用口径", level=1)
    add_callout(doc, "自动释放", "1 / 2 / 3 / 4 / 5 / 6 星依次解锁主动、被动、主动、被动、主动、被动。所有主动技能均为自动释放；没有合法目标时不施放，也不开始冷却。")
    p = doc.add_paragraph()
    set_paragraph(p, before=7, after=3)
    add_text(p, "统一战斗规则", size=11, color=DARK_BLUE, bold=True)
    for text in (
        "物理与法术伤害均使用施法者 attack，并走项目统一防御轴；伤害类型用于表现、抗性和后续扩展。",
        "当前目标区域最多 3 名单位，半径 260；当前一路最多 4 名单位；站位条件只读取本波战前格位。",
        "同一来源反复施加的护盾、攻击降低、减伤、标记等状态只刷新时间，除非技能明确写有层数。",
        "近战基准攻击距离为 300，远程基准攻击距离为 600；本批英雄均使用对应基准。",
    ):
        p = doc.add_paragraph(style="List Bullet")
        set_paragraph(p, after=4)
        add_text(p, text, size=10.2)

    doc.add_heading("归档结论", level=1)
    add_callout(doc, "结论", "本批 8 名英雄保持已审核通过的技能效果，不因本轮横向自检而改动。核心循环具备差异；后续落表与实测重点关注张辽/太史慈、马超/夏侯渊在标准队长战中的体验区隔。")
    p = doc.add_paragraph()
    set_paragraph(p, before=7, after=0)
    add_text(p, "历史例外记录：张辽、马超、荀彧的 6 星被动以其 5 星消耗满层资源为触发条件。该三项为已审核批准的例外，不作为后续新英雄的默认模板。", size=9.5, color=MUTED, italic=True)


def add_hero(doc, index, hero):
    doc.add_page_break()
    heading = doc.add_heading(hero["name"], level=1)
    heading.paragraph_format.keep_with_next = True

    p = doc.add_paragraph()
    set_paragraph(p, after=4)
    add_text(p, "定位：", size=10.5, color=DARK_BLUE, bold=True)
    add_text(p, f"{hero['battle']}；推荐{hero['position']}。", size=10.5)

    p = doc.add_paragraph()
    set_paragraph(p, after=4)
    add_text(p, "核心职责：", size=10.5, color=DARK_BLUE, bold=True)
    add_text(p, hero["role"], size=10.5)

    p = doc.add_paragraph()
    set_paragraph(p, after=4)
    add_text(p, "专属循环：", size=10.5, color=DARK_BLUE, bold=True)
    add_text(p, hero["core_name"], size=10.5)

    p = doc.add_paragraph()
    set_paragraph(p, after=8)
    add_text(p, hero["core"], size=10.5)

    p = doc.add_paragraph()
    set_paragraph(p, before=2, after=4)
    add_text(p, "六技能设计", size=11, color=DARK_BLUE, bold=True)
    add_skill_table(doc, hero["skills"])
    p = doc.add_paragraph()
    set_paragraph(p, before=8, after=0)
    add_text(p, "玩家可读性与 UI 提示：", size=9.5, color=DARK_BLUE, bold=True)
    add_text(p, hero["readability"], size=9.5, color=MUTED)


def add_closing(doc):
    doc.add_page_break()
    doc.add_heading("本批横向自检记录", level=1)
    p = doc.add_paragraph()
    set_paragraph(p, after=6)
    add_text(p, "自检结论：核心资源来源与玩家可见循环具有区分度，8 名英雄可保持当前设计归档；不进行整体返工。张辽/太史慈、马超/夏侯渊为后续实测对照组。", size=10.2)
    add_batch_matrix(doc)

    doc.add_heading("与已完成英雄的对比", level=2)
    add_history_comparison_table(doc)

    doc.add_page_break()
    doc.add_heading("落表提示", level=1)
    add_callout(doc, "配置边界", "本文件锁定的是已审核通过的技能逻辑、目标规则、触发条件与当前数值候选。统一落表时应为每个技能分配稳定 skill_id，并按技能逐项配置冷却、目标筛选、持续时间、内置冷却、状态叠加与可读性提示。")
    doc.add_heading("落表前实测项目", level=2)
    for text in (
        "自动释放条件：合法目标存在/不存在时的施放与冷却行为。",
        "标记与资源：震慑、战书、猎印、悬剑的刷新、转移、死亡清除、波次清空及每波上限。",
        "状态表现：护甲降低、攻击降低、护盾、攻速、减伤与对队长/精英增伤的刷新、叠加和结束时机。",
        "战斗场景：单一队长/精英、密集小兵、队长带兵、我方多名队长低血。",
    ):
        p = doc.add_paragraph(style="List Bullet")
        set_paragraph(p, after=4)
        add_text(p, text, size=10.2)

    doc.add_heading("归档依据", level=2)
    p = doc.add_paragraph()
    set_paragraph(p, after=0)
    add_text(p, "依据：第一批新增预选英雄技能设计候选稿 v0.1；Hero Skill Template Guideline；本批审核结论。本文不替代后续的正式英雄技能主源与逐技能数值配置表。", size=10, color=MUTED)


def main():
    doc = Document()
    configure_document(doc)
    add_overview(doc)
    for index, hero in enumerate(HEROES, start=1):
        add_hero(doc, index, hero)
    add_closing(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
