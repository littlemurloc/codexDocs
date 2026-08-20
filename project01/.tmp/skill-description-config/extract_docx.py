from pathlib import Path
from docx import Document


lines = []


def print_document(path_text: str) -> None:
    path = Path(path_text)
    document = Document(path)
    lines.append(f"# FILE: {path.name}")
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)
    for table_index, table in enumerate(document.tables, start=1):
        lines.append(f"## TABLE {table_index}")
        for row in table.rows:
            values = [cell.text.replace("\n", " / ").strip() for cell in row.cells]
            lines.append(" | ".join(values))


print_document(r"D:\codex\project01\deliverables\三国斗阵_Y3_英雄技能实现说明_v0.2.docx")
print_document(r"D:\codex\project01\deliverables\三国斗阵_第一批新增预选英雄技能设计归档_v1.0.docx")
Path(r"D:\codex\project01\.tmp\skill-description-config\source-docs.txt").write_text(
    "\n".join(lines), encoding="utf-8"
)
