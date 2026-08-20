from pathlib import Path
from docx import Document

path = Path(r"D:\codex\project01\deliverables\三国斗阵-军令效果设计清单-V1.docx")
document = Document(path)
lines = [f"# FILE: {path.name}"]

for paragraph in document.paragraphs:
    text = paragraph.text.strip()
    if text:
        lines.append(text)

for table_index, table in enumerate(document.tables, start=1):
    lines.append(f"## TABLE {table_index}")
    for row in table.rows:
        values = [cell.text.replace("\n", " / ").strip() for cell in row.cells]
        lines.append(" | ".join(values))

Path(r"D:\codex\project01\.tmp\skill-description-config\command-source-doc.txt").write_text(
    "\n".join(lines), encoding="utf-8"
)
