from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\codex\project01")
SOURCE = ROOT / "deliverables" / "第二批英雄技能设计候选稿-v0.2.md"
OUTPUT = ROOT / "deliverables" / "三国斗阵_第二批英雄技能设计阅读稿_v0.3.docx"

SHORT = {
    "许褚": [
        "对当前目标造成 190% 物理伤害，嘲讽 3 秒。",
        "普攻击中积累虎势；满 3 点震击附近最多 3 名敌人，各造成 90% 物理伤害。",
        "使附近最多 3 名敌人攻击降低 20%，持续 5 秒，并击退小兵。",
        "本队近战小兵受到的伤害降低 8%。",
        "获得 22% 最大生命护盾，持续 7 秒，并嘲讽附近最多 3 名敌人 3 秒。",
        "震地后，接下来 2 次普攻的攻击速度提高 30%。",
    ],
    "孙策": [
        "对当前目标造成 190% 物理伤害。",
        "接敌或主动技能命中积累潮势；满 2 点冲斩路径上最多 3 名敌人，各造成 110% 物理伤害。",
        "攻击当前一路最多 4 名敌人，各造成 120% 物理伤害，并击退小兵。",
        "潮势冲斩后受到的伤害降低 15%，持续 3 秒。",
        "对当前目标造成 300% 物理伤害，并打断其正在进行的可打断施法。",
        "潮势冲斩命中后返还 1 点潮势。",
    ],
    "孟获": [
        "对当前目标造成 200% 物理伤害，使其攻击降低 15%，持续 4 秒。",
        "本队近战小兵阵亡或孟获施放主动技能时积累蛮勇；满 3 点狂化 6 秒，攻速提高 20%，普攻额外横扫最多 2 名敌人，各造成 70% 物理伤害。",
        "孟获与本队存活小兵的攻击提高 15%，持续 6 秒。",
        "本队近战小兵受到的伤害降低 10%。",
        "对当前目标造成 310% 物理伤害，眩晕 1 秒。",
        "狂化期间每次普攻击中，延长狂化 0.5 秒，最多延长 3 秒。",
    ],
    "大乔": [
        "为生命比例最低的受伤友方队长恢复 14% 最大生命。",
        "我方队长施放主动技能后凝露；下一次普攻击中时消耗凝露，为生命比例最低的受伤队长恢复 5% 最大生命。",
        "对附近最多 3 名敌人各造成 120% 法术伤害，使其攻击降低 20%，持续 5 秒。",
        "受到凝露治疗的队长攻击提高 12%，持续 4 秒。",
        "友方队长生命低于 70% 时，为全体友方队长提供 12% 最大生命护盾，持续 6 秒。",
        "凝露普攻改为治疗生命比例最低的最多 2 名受伤队长，每人恢复 5% 最大生命。",
    ],
    "张郃": [
        "射击当前一路最前方敌人，造成 190% 物理伤害，减速 30%，持续 3 秒。",
        "普攻先射一路最前方敌人，下一箭改射最后方敌人，随后继续轮换。",
        "射击当前一路最后方敌人，造成 210% 物理伤害，使其攻击降低 20%，持续 5 秒。",
        "本队远程小兵造成的伤害提高 10%。",
        "分别射击一路最前方与最后方敌人，各造成 160% 物理伤害；只有一名敌人时两箭都命中它。",
        "后阵箭命中后，向该路最前方敌人补射一箭，造成 70% 物理伤害。",
    ],
    "祝融": [
        "对当前目标造成 180% 物理伤害，并种下炎种；炎种持续 5 秒，到期或目标死亡时爆炸，对附近最多 3 名敌人各造成 100% 物理伤害。",
        "普攻优先为未着火敌人种炎种；每次炎种爆炸获得 1 点火势，满 3 点释放火浪，对当前一路最多 4 名敌人各造成 130% 物理伤害。",
        "对附近最多 3 名敌人各造成 110% 物理伤害，并为未着火者种下炎种。",
        "带有炎种的敌人攻击降低 10%。",
        "对当前目标造成 220% 物理伤害，并立即引爆场上所有炎种。",
        "火浪命中后，为其中最多 2 名未着火敌人种下炎种。",
    ],
    "韩当": [
        "对当前目标造成 190% 物理伤害，使其攻击速度降低 25%，持续 4 秒。",
        "友方队长受击积累备矢；满 2 枚后援射伤害来源，造成 120% 物理伤害，并为受击队长提供 8% 最大生命护盾，持续 4 秒。",
        "对附近最多 3 名敌人各造成 125% 物理伤害，减速 30%，持续 3 秒。",
        "本队小兵受到的远程伤害降低 10%。",
        "对当前目标造成 280% 物理伤害，并为本队存活小兵提供 12% 最大生命护盾，持续 6 秒。",
        "援射命中后，使敌人攻击降低 20%，持续 5 秒。",
    ],
}


def read_heroes():
    heroes = []
    hero = None
    for line in SOURCE.read_text(encoding="utf-8").splitlines():
        match = re.match(r"^## (许褚|孙策|孟获|大乔|张郃|祝融|韩当)：(.+)$", line)
        if match:
            hero = {"name": match.group(1), "theme": match.group(2), "skills": []}
            heroes.append(hero)
            continue
        if hero is None:
            continue
        if line.startswith("**定位**："):
            hero["position"] = line.removeprefix("**定位**：").replace("**", "").strip()
        elif line.startswith("**核心机制："):
            hero["mechanic"] = line.replace("**", "")
        elif line.startswith("**一句话循环**："):
            hero["loop"] = line.removeprefix("**一句话循环**：").strip()
        elif re.match(r"^\| [1-6] \|", line):
            parts = [part.strip() for part in line.strip("|").split("|")]
            if len(parts) != 4:
                raise ValueError(f"Unexpected skill row: {line}")
            hero["skills"].append(parts)
    if [hero["name"] for hero in heroes] != list(SHORT):
        raise ValueError("Hero order or count changed")
    for hero in heroes:
        if len(hero["skills"]) != 6 or len(SHORT[hero["name"]]) != 6:
            raise ValueError(f"Wrong skill count: {hero['name']}")
        if any(int(row[0]) != i for i, row in enumerate(hero["skills"], start=1)):
            raise ValueError(f"Star sequence changed: {hero['name']}")
        if not all(key in hero for key in ("position", "mechanic", "loop")):
            raise ValueError(f"Missing hero summary: {hero['name']}")
    return heroes


def set_font(run, size, bold=False, color=RGBColor(0, 0, 0)):
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:color"), "D9D9D9")
        element.set(qn("w:sz"), "4")


def set_cell_margin(cell, top=75, left=85, bottom=75, right=85):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        tag = qn(f"w:{side}")
        element = margins.find(tag)
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_row_no_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def write_cell(cell, text, *, size=8.5, bold=False, white=False, center=False):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_border(cell)
    set_cell_margin(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(text), size, bold, RGBColor(255, 255, 255) if white else RGBColor(0, 0, 0))


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, before, after in (("Title", 18, 0, 8), ("Heading 1", 13, 14, 5), ("Heading 2", 10.5, 8, 4)):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        if name == "Title":
            p_pr = style._element.get_or_add_pPr()
            border = p_pr.find(qn("w:pBdr"))
            if border is not None:
                p_pr.remove(border)


def add_labeled_paragraph(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run(label), 9.5, True)
    set_font(p.add_run(value), 9.5)
    return p


def add_table(doc, headers, rows, widths, *, bold_second=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(round(sum(widths) * 1440)))
    for idx, header in enumerate(headers):
        table.columns[idx].width = Inches(widths[idx])
        table.rows[0].cells[idx].width = Inches(widths[idx])
        write_cell(table.rows[0].cells[idx], header, size=8.4, bold=True, white=True, center=idx < 2)
        set_cell_shading(table.rows[0].cells[idx], "39424A")
    set_repeat_header(table.rows[0])
    set_row_no_split(table.rows[0])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        set_row_no_split(table.rows[-1])
        for idx, value in enumerate(values):
            cells[idx].width = Inches(widths[idx])
            write_cell(cells[idx], value, size=8.35, bold=bold_second and idx == 1, center=idx == 0)
            if row_index % 2:
                set_cell_shading(cells[idx], "F7F9FA")
    return table


def build():
    heroes = read_heroes()
    doc = Document()
    style_document(doc)
    title = doc.add_paragraph("三国斗阵 第二批英雄技能设计阅读稿", style="Title")
    p_pr = title._p.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)
    p = doc.add_paragraph()
    set_font(p.add_run("候选版 v0.3  |  7 名英雄  |  42 项技能"), 9.5)
    doc.add_paragraph(
        "本稿将已逐人审核的第二批技能候选设计整理为便于阅读与核对的版本。每名英雄先说明战斗定位和核心循环，再并列展示正式效果与经审核的游戏内简短描述。简短描述用于玩家阅读建议，正式效果仍是制作依据。"
    )
    doc.add_paragraph(
        "本批尚未进入正式 HeroConfig：基础攻击类型、攻击距离、小兵类型及建议数值仍须配置审核和原型实测。文档不改变已审核的技能机制。"
    )

    doc.add_heading("使用口径", level=1)
    doc.add_paragraph(
        "1 至 6 星依次解锁主动、被动、主动、被动、主动、被动；已解锁主动技能自动释放。没有合法目标或未满足技能条件时不施放，也不开始冷却。技能描述中的伤害百分比均以英雄攻击为基数。"
    )
    doc.add_paragraph(
        "“当前目标附近”最多命中 3 名敌人，“当前一路”按目标所在路选敌；具体边界条件、叠加规则和数值仍以正式效果列及后续落表审核为准。"
    )

    doc.add_heading("英雄一览", level=1)
    add_table(
        doc,
        ["英雄", "战斗定位", "核心循环"],
        [(hero["name"], hero["position"].replace("；", "，"), hero["loop"]) for hero in heroes],
        [0.72, 2.28, 4.14],
    )

    for hero in heroes:
        heading = doc.add_heading(f"{hero['name']}  {hero['theme']}", level=1)
        heading.paragraph_format.page_break_before = True
        add_labeled_paragraph(doc, "定位：", hero["position"])
        add_labeled_paragraph(doc, "一句话循环：", hero["loop"])
        add_labeled_paragraph(doc, "核心机制：", hero["mechanic"].removeprefix("核心机制："))
        doc.add_heading("六技能设计", level=2)
        rows = []
        for i, (star, name, kind, formal) in enumerate(hero["skills"]):
            rows.append((f"{star} 星", f"{name}\n{kind}", formal, SHORT[hero["name"]][i]))
        add_table(
            doc,
            ["星级", "技能", "正式效果与自动规则", "游戏内描述建议"],
            rows,
            [0.5, 1.1, 2.78, 2.76],
            bold_second=True,
        )

    doc.save(OUTPUT)
    print(f"Created {OUTPUT}")
    print(f"Heroes: {len(heroes)}, skills: {sum(len(h['skills']) for h in heroes)}")


if __name__ == "__main__":
    build()
